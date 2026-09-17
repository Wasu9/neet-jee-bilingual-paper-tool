"""
NEET / JEE Simple DTP Tool v8.0
Master Page + 2 Column + Header/Footer + Page No + Word/PDF
"""

import streamlit as st
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
import re
from typing import List
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib import colors

st.set_page_config(page_title="Simple DTP Tool", page_icon="📄", layout="wide")

st.markdown("""
<style>
    .main-title { font-size: 1.6rem; font-weight: 700; color: #1a365d; }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">📄 Simple DTP Tool — NEET / JEE</div>', unsafe_allow_html=True)
st.caption("Master Page • 2 Column • Header/Footer • Page No • Word + PDF")

# ====================== HELPERS ======================
def set_cell_border(cell, color="999999", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), sz)
        tag.set(qn('w:color'), color)
        borders.append(tag)
    tcPr.append(borders)


def set_run_font(run, name="Times New Roman", size=10, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def split_blocks(text: str) -> List[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    parts = re.split(r'(?m)(?=^\s*(?:\d{1,3}|Q\.?\s*\d{1,3})[\.\)]\s+)', text)
    blocks = [p.strip() for p in parts if p.strip() and len(p.strip()) > 8]
    if len(blocks) <= 1:
        blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    return blocks


# ====================== WORD GENERATOR ======================
def create_word(
    institute: str,
    exam_name: str,
    class_name: str,
    subject: str,
    marks: str,
    date: str,
    instructions: str,
    eng_blocks: List[str],
    hin_blocks: List[str],
    margin_cm: float = 1.5,
    show_page_no: bool = True,
    header_text: str = "",
    footer_text: str = "",
) -> bytes:

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(margin_cm)
    sec.bottom_margin = Cm(margin_cm)
    sec.left_margin = Cm(margin_cm)
    sec.right_margin = Cm(margin_cm)

    # Header
    if header_text.strip():
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(header_text)
        set_run_font(run, "Arial", 9, True)

    # Footer + Page Number
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if footer_text.strip():
        run = fp.add_run(footer_text + "  |  ")
        set_run_font(run, "Arial", 8)

    if show_page_no:
        run = fp.add_run("Page ")
        set_run_font(run, "Arial", 8)
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run2 = fp.add_run()
        run2._r.append(fldChar1)
        run2._r.append(instrText)
        run2._r.append(fldChar2)
        set_run_font(run2, "Arial", 8)

    # Institute
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(institute)
    set_run_font(run, "Arial", 14, True, (26, 54, 93))

    # Exam name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(exam_name)
    set_run_font(run, "Arial", 12, True)

    # Info line
    info = f"Class: {class_name}   |   Subject: {subject}   |   Marks: {marks}   |   Date: {date}"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(info)
    set_run_font(run, "Arial", 9)

    # Instructions
    if instructions.strip():
        p = doc.add_paragraph()
        run = p.add_run("Instructions:")
        set_run_font(run, "Arial", 9, True)
        p = doc.add_paragraph()
        run = p.add_run(instructions)
        set_run_font(run, "Arial", 8)
        p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)

    # Questions (2 Column)
    max_len = max(len(eng_blocks), len(hin_blocks))
    eng_blocks += [""] * (max_len - len(eng_blocks))
    hin_blocks += [""] * (max_len - len(hin_blocks))

    for eng, hin in zip(eng_blocks, hin_blocks):
        if not eng.strip() and not hin.strip():
            continue

        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Cm(8.7)
        table.columns[1].width = Cm(8.7)

        for cell in table.rows[0].cells:
            set_cell_border(cell, "AAAAAA", "4")

        c0 = table.rows[0].cells[0]
        p0 = c0.paragraphs[0]
        run0 = p0.add_run(eng)
        set_run_font(run0, "Times New Roman", 9)
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)

        c1 = table.rows[0].cells[1]
        p1 = c1.paragraphs[0]
        run1 = p1.add_run(hin)
        set_run_font(run1, "Mangal", 9)
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)

        gap = doc.add_paragraph()
        gap.paragraph_format.space_before = Pt(3)
        gap.paragraph_format.space_after = Pt(3)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ====================== PDF GENERATOR ======================
def create_pdf(
    institute: str,
    exam_name: str,
    class_name: str,
    subject: str,
    marks: str,
    date: str,
    instructions: str,
    eng_blocks: List[str],
    hin_blocks: List[str],
    margin_cm: float = 1.5,
) -> bytes:

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=margin_cm*cm,
        rightMargin=margin_cm*cm,
        topMargin=margin_cm*cm,
        bottomMargin=margin_cm*cm
    )

    styles = getSampleStyleSheet()
    style_title = ParagraphStyle('Title', parent=styles['Normal'], fontSize=13, alignment=TA_CENTER, spaceAfter=4, fontName='Helvetica-Bold')
    style_sub = ParagraphStyle('Sub', parent=styles['Normal'], fontSize=10, alignment=TA_CENTER, spaceAfter=3, fontName='Helvetica-Bold')
    style_info = ParagraphStyle('Info', parent=styles['Normal'], fontSize=8, alignment=TA_CENTER, spaceAfter=6)
    style_inst = ParagraphStyle('Inst', parent=styles['Normal'], fontSize=8, alignment=TA_LEFT, spaceAfter=8)
    style_q = ParagraphStyle('Q', parent=styles['Normal'], fontSize=8, alignment=TA_LEFT, leading=11)

    story = []

    story.append(Paragraph(institute, style_title))
    story.append(Paragraph(exam_name, style_sub))
    story.append(Paragraph(f"Class: {class_name} | Subject: {subject} | Marks: {marks} | Date: {date}", style_info))

    if instructions.strip():
        story.append(Paragraph(f"<b>Instructions:</b> {instructions}", style_inst))

    story.append(Spacer(1, 6))

    max_len = max(len(eng_blocks), len(hin_blocks))
    eng_blocks += [""] * (max_len - len(eng_blocks))
    hin_blocks += [""] * (max_len - len(hin_blocks))

    for eng, hin in zip(eng_blocks, hin_blocks):
        if not eng.strip() and not hin.strip():
            continue
        data = [[Paragraph(eng.replace("\n", "<br/>"), style_q),
                 Paragraph(hin.replace("\n", "<br/>"), style_q)]]
        t = Table(data, colWidths=[8.5*cm, 8.5*cm])
        t.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t)
        story.append(Spacer(1, 4))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


# ====================== UI ======================
st.markdown("### 1. Master Page (Institute Details)")

c1, c2 = st.columns(2)
with c1:
    institute = st.text_input("Institute Name", value="SHAHEEN GROUP OF INSTITUTIONS")
    exam_name = st.text_input("Exam Name", value="GRAND EXAMINATION")
    class_name = st.text_input("Class", value="I PUC JEE")
with c2:
    subject = st.text_input("Subject", value="PCM")
    marks = st.text_input("Marks", value="300")
    date = st.text_input("Date", value="16/09/2026")

instructions = st.text_area("Instructions (optional)", height=70, placeholder="1. Calculator not allowed\n2. ...")

st.markdown("### 2. Page Settings")
c3, c4, c5 = st.columns(3)
with c3:
    margin = st.slider("Margin (cm)", 1.0, 2.5, 1.5, 0.1)
with c4:
    show_page_no = st.checkbox("Show Page Number", value=True)
with c5:
    header_text = st.text_input("Header Text", value="")
footer_text = st.text_input("Footer Text", value="")

st.markdown("### 3. Content (English | Hindi)")

col_eng, col_hin = st.columns(2)
with col_eng:
    st.markdown("**English Questions**")
    eng_text = st.text_area("eng", height=350, label_visibility="collapsed",
                            placeholder="1. Question...\n(1) option\n(2) option\n\n2. Next question...")
with col_hin:
    st.markdown("**Hindi Questions**")
    hin_text = st.text_area("hin", height=350, label_visibility="collapsed",
                            placeholder="1. प्रश्न...\n(1) विकल्प\n(2) विकल्प\n\n2. अगला प्रश्न...")

st.markdown("---")

if st.button("🚀 Generate Word + PDF", type="primary", use_container_width=True):
    eng_blocks = split_blocks(eng_text)
    hin_blocks = split_blocks(hin_text)

    if not eng_blocks and not hin_blocks:
        st.warning("Please paste some content.")
    else:
        with st.spinner("Generating..."):
            word_bytes = create_word(
                institute, exam_name, class_name, subject, marks, date, instructions,
                eng_blocks, hin_blocks, margin, show_page_no, header_text, footer_text
            )
            pdf_bytes = create_pdf(
                institute, exam_name, class_name, subject, marks, date, instructions,
                eng_blocks, hin_blocks, margin
            )

        st.success(f"✅ Generated • {max(len(eng_blocks), len(hin_blocks))} blocks")

        c1, c2 = st.columns(2)
        with c1:
            st.download_button("📥 Download Word", data=word_bytes,
                               file_name="bilingual_dtp.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)
        with c2:
            st.download_button("📥 Download PDF", data=pdf_bytes,
                               file_name="bilingual_dtp.pdf",
                               mime="application/pdf",
                               use_container_width=True)

st.markdown("---")
st.caption("v8.0 Simple DTP • Master Page • 2 Column • Header/Footer • Page No • Word + PDF")
